"""Бот аудиту дзвінків (локальний режим).

Завантажує аудіо та таблицю з Google Drive на ПК, транскрибує (Whisper),
аналізує роботу менеджера (OpenAI) і зберігає заповнену оцінкову таблицю
локально як .xlsx.
"""
import datetime
import os
import re

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

import analyze
import config
import transcribe
from google_client import XLSX_MIME, GoogleClient

RED_FILL = PatternFill("solid", fgColor="F4CCCC")
SCORE_COLUMN_TITLE = "Бали"

# Ключові слова для пошуку рядка-шапки та ролей колонок.
HEADER_KEYWORDS = ("дата", "тип", "номер", "телефон", "менеджер", "оцінк",
                   "коментар", "результат", "запис", "розмов", "філія")

# Розбір імені файлу виду 2025-09-10_15-52_0632838007_incoming.mp3
FILENAME_RE = re.compile(r"(\d{4}-\d{2}-\d{2}).*?(\d{7,15})")


def parse_filename(name):
    """Витягти дату та номер телефону з імені аудіофайлу."""
    m = FILENAME_RE.search(name)
    call_date, phone = None, ""
    if m:
        try:
            call_date = datetime.datetime.strptime(m.group(1), "%Y-%m-%d")
        except ValueError:
            call_date = None
        phone = m.group(2)
    return call_date, phone


def parse_date(value):
    """Розбір рядка дати у datetime. Повертає (datetime|None, has_time)."""
    s = str(value).strip()
    for fmt, has_time in (("%Y-%m-%d %H:%M", True), ("%Y-%m-%d", False)):
        try:
            return datetime.datetime.strptime(s[:16 if has_time else 10], fmt), has_time
        except (ValueError, TypeError):
            continue
    return None, False


def detect_header_row(rows):
    """Індекс рядка-шапки — рядок із найбільшою кількістю збігів ключових слів."""
    best_idx, best_hits = 0, -1
    for idx, row in enumerate(rows[:12]):
        hits = sum(
            any(kw in str(cell).lower() for kw in HEADER_KEYWORDS)
            for cell in row
            if cell
        )
        if hits > best_hits:
            best_idx, best_hits = idx, hits
    return best_idx


def column_role(name):
    """Роль колонки за її назвою (None — звичайна колонка для заповнення)."""
    n = name.lower()
    if "дата" in n and "запис" not in n:
        return "date"
    if "телефон" in n or "номер" in n:
        return "phone"
    if "філія" in n:
        return "branch"
    if "менеджер" in n:
        return "manager"
    if "оцінк" in n:
        return "score"
    if "коментар" in n:
        return "comment"
    if "яка робота" in n or "робота з топ" in n:
        return "choice"
    return None


def parse_inline_list(formula1):
    """Розбір inline-списку Excel data validation у перелік варіантів.

    Довгі списки Excel зберігає як склейку рядків ("a,b"&"c,d") — її знімаємо.
    Повертає список варіантів або None (для діапазонних / #REF! формул).
    """
    f = (formula1 or "").strip()
    if not (f.startswith('"') and f.endswith('"')):
        return None
    f = f.replace('"&"', "")  # склейка довгих списків
    return [o.strip() for o in f[1:-1].split(",") if o.strip()]


def parse_validations(ws):
    """Зчитати випадаючі списки таблиці.

    Повертає {col_index: {"kind": "binary"|"choice", "options": [...]}}.
    """
    info = {}
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list":
            continue
        options = parse_inline_list(dv.formula1)
        if not options:
            continue
        kind = "binary" if {o.lower() for o in options} <= {"0", "1"} else "choice"
        for rng in dv.sqref.ranges:
            for col in range(rng.min_col, rng.max_col + 1):
                info[col - 1] = {"kind": kind, "options": options}
    return info


def extend_validations(ws, new_rows):
    """Поширити наявні випадаючі списки таблиці на нові рядки."""
    for dv in ws.data_validations.dataValidation:
        cols = set()
        for rng in dv.sqref.ranges:
            cols.update(range(rng.min_col, rng.max_col + 1))
        for r in new_rows:
            for c in cols:
                dv.add(f"{get_column_letter(c)}{r}")


def build_layout(rows, ws):
    """Аналіз структури таблиці. Повертає словник з описом колонок.

    Тип і варіанти колонок беруться з випадаючих списків таблиці (data
    validations) — це найнадійніше джерело. Якщо їх немає, тип визначає модель.
    """
    header_idx = detect_header_row(rows)
    header = rows[header_idx]

    named = [
        {"index": idx, "name": str(raw).strip()}
        for idx, raw in enumerate(header)
        if raw not in (None, "")
    ]
    validations = parse_validations(ws)
    has_binary = any(v["kind"] == "binary" for v in validations.values())

    # Резервна класифікація моделлю — лише якщо в таблиці немає списків 1/0.
    kinds = {}
    if not has_binary:
        group_headers = ([str(c).strip() for c in rows[0] if c]
                         if header_idx > 0 else [])
        kinds = analyze.classify_columns(named, group_headers)

    columns = []  # {index, name, role, kind, options}
    roles_taken = set()
    for col in named:
        idx, name = col["index"], col["name"]
        role = column_role(name)
        if role in roles_taken:  # роль присвоюється лише першій колонці
            role = None
        if role:
            roles_taken.add(role)

        options = None
        if role in ("date", "phone", "branch"):
            kind = role
        elif idx in validations:
            kind = validations[idx]["kind"]
            options = validations[idx]["options"]
        elif "запис" in name.lower() and "дата" in name.lower():
            kind = "date"  # колонка дати запису на сервіс
        elif role == "score":
            kind = "score"
        elif role == "choice":
            kind = "choice"
        else:
            kind = kinds.get(idx, "text")
        columns.append({"index": idx, "name": name, "role": role,
                        "kind": kind, "options": options})

    return {
        "header_idx": header_idx,
        "columns": columns,
        "score_col_idx": max((c["index"] for c in columns), default=-1) + 1,
    }


def main():
    print("=== Бот аудиту дзвінків (локальний режим) ===")
    gc = GoogleClient()

    # --- Крок 1: завантаження аудіо на ПК ---
    print("\n[1] Пошук аудіофайлів у вихідній папці...")
    source_audio = gc.list_audio(config.SOURCE_FOLDER_ID)
    if config.MAX_FILES > 0:
        source_audio = source_audio[: config.MAX_FILES]
    print(f"    Знайдено {len(source_audio)} аудіофайл(ів).")

    calls = []
    for f in source_audio:
        local_audio = os.path.join(config.WORK_DIR, f["name"])
        if os.path.exists(local_audio):
            print(f"    ~ вже завантажено: {f['name']}")
        else:
            gc.download_file(f["id"], local_audio)
            print(f"    + завантажено: {f['name']}")
        calls.append({"name": f["name"], "audio": local_audio})

    if not calls:
        print("\nНемає аудіофайлів для обробки. Завершено.")
        return

    # --- Крок 2: транскрипція ---
    print("\n[2] Транскрипція аудіофайлів...")
    for call in calls:
        base = os.path.splitext(call["name"])[0]
        local_txt = os.path.join(config.WORK_DIR, base + ".txt")
        if os.path.exists(local_txt):
            with open(local_txt, encoding="utf-8") as fh:
                call["transcript"] = fh.read()
            print(f"    ~ транскрипт існує: {base}.txt")
        else:
            text = transcribe.transcribe(call["audio"])
            with open(local_txt, "w", encoding="utf-8") as fh:
                fh.write(text)
            call["transcript"] = text
            print(f"    + транскрибовано: {call['name']}")
        call["txt"] = local_txt

    # --- Крок 3: завантаження таблиці-шаблону ---
    print("\n[3] Завантаження оцінкової таблиці...")
    result_name = "Звіт прослуханих розмов.xlsx"
    result_path = os.path.join(config.WORK_DIR, result_name)
    gc.export_spreadsheet(config.SOURCE_SHEET_ID, result_path)

    wb = load_workbook(result_path)
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    layout = build_layout(rows, ws)
    columns = layout["columns"]
    header_row = layout["header_idx"] + 1  # 1-індекс
    score_idx = layout["score_col_idx"]

    by_role = {c["role"]: c for c in columns if c["role"]}
    fillable = [c for c in columns
                if c["role"] not in ("date", "phone", "branch")]
    binary_cols = [c for c in fillable if c["kind"] == "binary"]
    choice_cols = [c for c in fillable if c["kind"] == "choice"]
    comment_col = by_role.get("comment")

    print(f"    Шапка: рядок {header_row}; критеріїв 1/0: {len(binary_cols)}; "
          f"колонок з випадаючим списком: {len(choice_cols)}")

    # Заголовок колонки балів.
    ws.cell(row=header_row, column=score_idx + 1, value=SCORE_COLUMN_TITLE)

    # --- Крок 4: аналіз та заповнення рядків ---
    print("\n[4] Аналіз дзвінків та заповнення таблиці...")
    next_row = ws.max_row + 1
    negative = 0
    schema = [{"name": c["name"], "kind": c["kind"], "options": c["options"]}
              for c in fillable]

    for i, call in enumerate(calls):
        excel_row = next_row + i

        # Дата та телефон — з імені файлу.
        call_date, phone = parse_filename(call["name"])
        call_date_iso = call_date.strftime("%Y-%m-%d") if call_date else None

        result = analyze.analyze(call["transcript"], schema, call_date_iso)
        values = result["values"]

        if by_role.get("date") and call_date:
            cell = ws.cell(row=excel_row, column=by_role["date"]["index"] + 1,
                           value=call_date)
            cell.number_format = "YYYY-MM-DD"
        if by_role.get("phone") and phone:
            ws.cell(row=excel_row, column=by_role["phone"]["index"] + 1, value=phone)

        # Колонки, заповнені моделлю.
        total = 0
        for col in fillable:
            val = values.get(col["name"], "")
            cell = ws.cell(row=excel_row, column=col["index"] + 1)
            if col["kind"] == "date" and val:
                parsed, has_time = parse_date(val)
                if parsed:
                    cell.value = parsed
                    cell.number_format = "YYYY-MM-DD HH:MM" if has_time else "YYYY-MM-DD"
                else:
                    cell.value = val
            else:
                cell.value = val
            if col["kind"] == "binary":
                total += int(val or 0)
        ws.cell(row=excel_row, column=score_idx + 1, value=total)

        # Червона заливка негативного коментаря.
        if result["comment_is_negative"] and comment_col:
            ws.cell(row=excel_row,
                    column=comment_col["index"] + 1).fill = RED_FILL
            negative += 1

        flag = "  ⚠ НЕ ОК" if result["comment_is_negative"] else ""
        print(f"    [{i + 1}/{len(calls)}] {call['name']}: бали={total}{flag}")

    # Поширити випадаючі списки (тип робіт, критерії 1/0 тощо) на нові рядки.
    extend_validations(ws, range(next_row, next_row + len(calls)))

    # --- Крок 5: збереження результату ---
    wb.save(result_path)
    print(f"\n[5] Збережено локально: {result_path}")

    # --- Крок 6: вивантаження на Google Drive ---
    print("\n[6] Вивантаження результатів на Google Drive...")
    for call in calls:
        gc.upload(call["txt"], config.SOURCE_FOLDER_ID, mime_type="text/plain")
        print(f"    + транскрипт: {os.path.basename(call['txt'])}")
    gc.upload(result_path, config.REPORT_FOLDER_ID, mime_type=XLSX_MIME)
    print(f"    + звіт: {result_name}")

    print(f"\n[OK] Готово. Внесено {len(calls)} дзвінків.")
    print(f"     Негативних коментарів (червоні): {negative}")


if __name__ == "__main__":
    main()
